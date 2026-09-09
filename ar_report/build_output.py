"""Assembles the final 'AR Aging' workbook from the cleaned/aggregated data.

Deliberately does NOT reuse the team's original workbook as a base (openpyxl's
round-tripping of PivotTable objects is unreliable, and re-saving a 23MB file
that contains four live pivots risked silently corrupting them). Instead this
builds a fresh workbook sheet-by-sheet with plain values -- the four pivots'
*numbers* are already correct (computed in aggregate.py), so nothing here
needs to be a live, refreshable Excel PivotTable to be correct.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Color, Font, GradientFill, PatternFill
from openpyxl.styles.fills import Stop
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Colors below are read directly from the real workbook (via Excel COM,
# not guessed) -- the team's original header styling per sheet, not an
# invented scheme. Most sheets use no header fill at all (plain bold black
# text); only AR Provision New's "blue" helper columns, By Invoice, and By
# Customer have an explicit header color in the original.
DEFAULT_HEADER_FILL = None
DEFAULT_HEADER_FONT = Font(color="000000", bold=True)
HELPER_FILL = PatternFill("solid", fgColor="156082", bgColor="156082")
HELPER_FONT = Font(color="CAE5FB", bold=True)
BY_INVOICE_HEADER_FILL = PatternFill("solid", fgColor="FFFFCC", bgColor="FFFFCC")
BY_INVOICE_HEADER_FONT = Font(color="000000", bold=True)
BY_CUSTOMER_HEADER_FILL = PatternFill("solid", fgColor="C0E6F5", bgColor="C0E6F5")
BY_CUSTOMER_HEADER_FONT = Font(color="000000", bold=True)
MONTH_END_HEADER_FILL = PatternFill("solid", fgColor="A02B93", bgColor="A02B93")
MONTH_END_HEADER_FONT = Font(color="FFFFFF", bold=True)
# Row 1 of By Customer / By Invoice: navy band across the full row, with a
# SUBTOTAL(9, ...) per money column so the totals follow the autofilter.
TOTALS_FILL = PatternFill("solid", fgColor="0E2841", bgColor="0E2841")
TOTALS_FONT = Font(color="FFFFFF", bold=True)
BY_CUSTOMER_TOTALS_FORMAT = '_-[$$-409]* #,##0_ ;_-[$$-409]* -#,##0 ;_-[$$-409]* "-"??_ ;_-@_ '
BY_INVOICE_TOTALS_FORMAT = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
# The real workbook's conditional formats are radial gradients -- white at
# the cell centre fading out to colour at the edges -- not flat solid fills.
# Colours and geometry are read straight from its dxf records.
def _radial_fill(edge_rgb: str) -> GradientFill:
    return GradientFill(
        type="path", left=0.5, right=0.5, top=0.5, bottom=0.5,
        stop=[Stop(Color(theme=0), 0), Stop(Color(rgb=edge_rgb), 1)],
    )


OVERDUE_FILL_1 = _radial_fill("FFF8F8BA")  # Ageing 199-239: light yellow
OVERDUE_FILL_2 = _radial_fill("FFFFBDBD")  # Ageing >239 and overdue Ar Balance: light pink

CURRENCY_FORMAT = "#,##0.00"
DATE_FORMAT = "dd/mm/yyyy"


def _write_df(
    ws: Worksheet,
    df: pd.DataFrame,
    start_row: int = 1,
    start_col: int = 1,
    header_styles: dict[str, tuple[PatternFill, Font]] | None = None,
    currency_cols: set[str] | None = None,
    date_cols: set[str] | None = None,
    header_fill: PatternFill | None = DEFAULT_HEADER_FILL,
    header_font: Font = DEFAULT_HEADER_FONT,
) -> int:
    header_styles = header_styles or {}
    currency_cols = currency_cols or set()
    date_cols = date_cols or set()

    for col_offset, col_name in enumerate(df.columns):
        col_idx = start_col + col_offset
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        if col_name in header_styles:
            cell.fill, cell.font = header_styles[col_name]
        else:
            if header_fill is not None:
                cell.fill = header_fill
            cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_offset, (col_name, value) in enumerate(zip(df.columns, row)):
            col_idx = start_col + col_offset
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_name in currency_cols:
                cell.number_format = CURRENCY_FORMAT
            elif col_name in date_cols:
                cell.number_format = DATE_FORMAT

    last_row = start_row + len(df)
    if start_col == 1:
        last_col = len(df.columns)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(last_col)}{last_row}"
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for col_offset, col_name in enumerate(df.columns):
        col_idx = start_col + col_offset
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(28, len(str(col_name)) + 2))
    return last_row


def _write_totals_row(
    ws: Worksheet,
    df: pd.DataFrame,
    row: int,
    total_cols: set[str],
    first_data_row: int,
    last_row: int,
    number_format: str,
) -> None:
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = TOTALS_FILL
        cell.font = TOTALS_FONT
        if col_name in total_cols:
            letter = get_column_letter(col_idx)
            cell.value = f"=SUBTOTAL(9,{letter}{first_data_row}:{letter}{last_row})"
            cell.number_format = number_format


AR_PROVISION_HELPER_COLS = {
    "Ageing",
    "Overdue days",
    "Region",
    "Ar Balance2",
    "Aging Bracket",
    "Updated Status",
    "Additional due End of month",
    "Invoice Value",
    "On Account2",
    "Not Due (helper)",
    "Aging 1 to 30 (calc)",
    "Aging 31 to 60 (calc)",
    "Aging 61 to 90 (calc)",
    "Aging 91 to 120 (calc)",
    "Aging 121 to 150 (calc)",
    "Aging >=151 (calc)",
}


def build_workbook(
    ar_df: pd.DataFrame,
    pdc_df: pd.DataFrame,
    insurance_df: pd.DataFrame,
    backlog_df: pd.DataFrame,
    as_on_date: pd.Timestamp,
    backlog_detail_pivot: pd.DataFrame,
    backlog_by_customer: pd.DataFrame,
    by_customer: pd.DataFrame,
    by_invoice: pd.DataFrame,
    lookups,
) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("AR Provision New")
    ws.cell(row=1, column=1, value="As on Date:").font = Font(bold=True)
    date_cell = ws.cell(row=1, column=2, value=pd.Timestamp(as_on_date).to_pydatetime())
    date_cell.number_format = DATE_FORMAT
    _write_df(
        ws,
        ar_df,
        start_row=3,
        header_styles={c: (HELPER_FILL, HELPER_FONT) for c in AR_PROVISION_HELPER_COLS},
        currency_cols={
            "Ar Balance", "Ar Balance2", "On Account", "Not Due Amount", "Invoice Value",
            "Additional due End of month", "Total Insurance Limit",
        },
        date_cols={"Document Date", "Document Due Date", "As Of Dt"},
    )

    ws = wb.create_sheet("PDC list as per Orion")
    _write_df(ws, pdc_df, currency_cols={"Cheque Amount", "LC Amount"}, date_cols={"Due Dt", "Doc Dt"})

    ws = wb.create_sheet("Look up links")
    ws.cell(row=1, column=1, value="Insurance Report (pasted from ORION)").font = Font(bold=True)
    _write_df(ws, insurance_df, start_row=2, currency_cols={"Insurance Limit"})

    trading_start_col = len(insurance_df.columns) + 2
    ws.cell(row=1, column=trading_start_col, value="Trading Experience -- DSO vs Credit Terms").font = Font(bold=True)
    _write_df(ws, lookups.trading_df, start_row=2, start_col=trading_start_col)

    ws = wb.create_sheet("Backlog")
    _write_df(ws, backlog_df, currency_cols={"Pending Val (Lc)"}, date_cols={"Order Date", "Order Delivery Date"})
    pivot_start_col = len(backlog_df.columns) + 2
    for col_offset, col_name in enumerate(backlog_detail_pivot.columns):
        c = ws.cell(row=1, column=pivot_start_col + col_offset, value=col_name)
        c.font = DEFAULT_HEADER_FONT
    for row_idx, row in enumerate(backlog_detail_pivot.itertuples(index=False), start=2):
        for col_offset, value in enumerate(row):
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            ws.cell(row=row_idx, column=pivot_start_col + col_offset, value=value)

    ws = wb.create_sheet("By Customer")
    money_cols = set(by_customer.columns) - {"Cust Code", "Cust Name", "Main Ac", "EZ#", "Country", "Region", "CustStatus", "CT", "DSO"}
    month_end_cols = {"Remaining Month End Dues", "TOTAL  Overdues till Month End"}
    last_row = _write_df(
        ws, by_customer, start_row=2, currency_cols=money_cols,
        header_fill=BY_CUSTOMER_HEADER_FILL, header_font=BY_CUSTOMER_HEADER_FONT,
        header_styles={c: (MONTH_END_HEADER_FILL, MONTH_END_HEADER_FONT) for c in month_end_cols},
    )
    _write_totals_row(
        ws, by_customer, row=1, total_cols=money_cols - {"Insured Limit"},
        first_data_row=3, last_row=last_row, number_format=BY_CUSTOMER_TOTALS_FORMAT,
    )

    ws = wb.create_sheet("By Invoice")
    money_cols = {"On Account", "Not Due", "Ar Balance", "Aging 1 to 30", "Aging 31 to 60", "Aging 61 to 90", "Aging 91 to 120", "Aging 121 to 150", "Aging >=151", "Total Insurance Limit", "LC & BG Guarantee"}
    last_row = _write_df(
        ws, by_invoice, start_row=2, currency_cols=money_cols, date_cols={"Document Date", "Document Due Date"},
        header_fill=BY_INVOICE_HEADER_FILL, header_font=BY_INVOICE_HEADER_FONT,
    )
    _write_totals_row(
        ws, by_invoice, row=1, total_cols=money_cols - {"Total Insurance Limit", "LC & BG Guarantee"},
        first_data_row=3, last_row=last_row, number_format=BY_INVOICE_TOTALS_FORMAT,
    )

    ageing_col_letter = get_column_letter(list(by_invoice.columns).index("Ageing") + 1)
    ar_balance_col_letter = get_column_letter(list(by_invoice.columns).index("Ar Balance") + 1)
    due_date_col_letter = get_column_letter(list(by_invoice.columns).index("Document Due Date") + 1)
    data_range = f"{ageing_col_letter}3:{ageing_col_letter}{last_row}"
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator="between", formula=["199", "239"], fill=OVERDUE_FILL_1),
    )
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator="greaterThan", formula=["239"], fill=OVERDUE_FILL_2),
    )
    ar_balance_range = f"{ar_balance_col_letter}3:{ar_balance_col_letter}{last_row}"
    ws.conditional_formatting.add(
        ar_balance_range,
        FormulaRule(
            formula=[f"${due_date_col_letter}3<='AR Provision New'!$B$1"],
            fill=OVERDUE_FILL_2,
        ),
    )

    ws = wb.create_sheet("Look up links backlog pivot")
    _write_df(ws, backlog_by_customer, currency_cols={"Sum of Pending Val (Lc)"})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
